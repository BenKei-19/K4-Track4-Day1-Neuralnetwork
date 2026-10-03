"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
           "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
           "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
           "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
MAX_ROWS = 60   # mẫu có sẵn công thức cho dòng 2..61

# tên trong code -> tên trong danh sách chọn (data validation) của mẫu
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(x):
    """NaN/inf -> None (JSON chuẩn và ô Excel trống); áp dụng đệ quy cho dict/list."""
    if isinstance(x, float) and not math.isfinite(x):
        return None
    if isinstance(x, dict):
        return {k: _clean(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_clean(v) for v in x]
    return x


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = {k: result[k] for k in ("cfg", "history", "summary")}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(_clean(payload), f, ensure_ascii=False, indent=1)
    return str(path)


def load_result(path: str) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir (không đọc thư mục con), trả về danh sách dict (sắp theo exp_id)."""
    return sorted((load_result(p) for p in Path(results_dir).glob("*.json")), key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng (số do scripts/evaluate.py in ra)."""
    cfg, s = result["cfg"], result["summary"]
    row = {
        "exp_id": cfg["exp_id"], "group": cfg["group"], "description": cfg["description"],
        "loss": LOSS_NAMES[cfg["loss"]], "optimizer": OPT_NAMES[cfg["optimizer"]],
        "lr": cfg["lr"], "weight_decay": cfg["weight_decay"], "batch": cfg["batch"], "epochs": cfg["epochs"],
        "hidden": "-".join(str(h) for h in cfg["hidden"]), "dropout": cfg["dropout"],
        "clip_norm": "none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        "precision": cfg["precision"], "init": cfg["init"], "seed": cfg["seed"],
        "diverged": "Y" if s["diverged"] else "N",
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes,
    }
    for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
              "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        v = _clean(s.get(k))
        row[k] = round(v, 6) if isinstance(v, float) else v
    if eval_scores is not None:
        row["eval_acc"] = round(eval_scores["accuracy"], 6)
        row["eval_macro_f1"] = round(eval_scores["macro_f1"], 6)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seeds: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    - Đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào; BỎ QUA các cột công thức.
    - Xoá giá trị mẫu cũ ở các cột nhập liệu (dòng baseline ví dụ của mẫu) trước khi ghi.
    - seeds: danh sách exp_id baseline khác seed -> cột A của sheet Seeds (dòng 2..6).
    - summary_notes: {group: nhận xét} -> cột "nhận xét ngắn" của sheet Summary.
    Không dùng data_only=True (sẽ mất công thức). Mở file bằng Excel/LibreOffice để công thức tính lại.
    """
    import openpyxl

    assert len(rows) <= MAX_ROWS, f"mẫu chỉ có công thức cho {MAX_ROWS} dòng"
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id phải duy nhất"
    wb = openpyxl.load_workbook(template_path)

    ws = wb["Experiments"]
    header = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    missing = [c for c in COLUMNS if c not in header]
    assert not missing, f"mẫu thiếu cột {missing}"
    for r in range(2, MAX_ROWS + 2):
        for col in COLUMNS:
            ws.cell(r, header[col]).value = None
    for i, row in enumerate(rows):
        for col in COLUMNS:
            v = row.get(col)
            if v is not None and v != "":
                ws.cell(i + 2, header[col]).value = v

    if seeds is not None:
        ws_seed = wb["Seeds"]
        assert len(seeds) <= 5, "sheet Seeds có 5 ô (A2:A6)"
        for r in range(2, 7):
            ws_seed.cell(r, 1).value = seeds[r - 2] if r - 2 < len(seeds) else None

    if summary_notes:
        ws_sum = wb["Summary"]
        note_col = next(c for c in range(1, ws_sum.max_column + 1)
                        if str(ws_sum.cell(1, c).value or "").startswith("nhận xét"))
        for r in range(2, ws_sum.max_row + 1):
            g = ws_sum.cell(r, 1).value
            if g in summary_notes:
                ws_sum.cell(r, note_col).value = summary_notes[g]

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
